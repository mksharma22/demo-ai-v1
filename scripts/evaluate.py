"""Evaluate decisions against human labels and sweep thresholds.

Input: .xlsx with AddressA, AddressB and TrueLabel (SAME / DIFFERENT / UNKNOWN).
Rows labelled UNKNOWN (or unlabelled) are excluded from precision/recall.

Definitions (n_same = rows labelled SAME, n_diff = rows labelled DIFFERENT):
  precision         = MATCH&SAME / MATCH                (of auto-matches, how many are right)
  recall            = MATCH&SAME / n_same               (REVIEW counts as not matched)
  false_match_rate  = MATCH&DIFFERENT / n_diff          (most costly error in most uses)
  false_non_match   = NON-MATCH&SAME / n_same
  review_rate       = REVIEW / labelled rows            (manual workload)

Usage:
  python scripts/evaluate.py data/sample/sample_input.xlsx
  python scripts/evaluate.py labelled.xlsx --sweep
"""

from __future__ import annotations

import argparse
import dataclasses
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from addrmatch.matcher import AddressMatcher  # noqa: E402
from addrmatch.normalization import CachingNormalizer, get_libpostal_backend  # noqa: E402
from addrmatch.scoring_config import ScoringConfig, load_scoring_config  # noqa: E402
from addrmatch.settings import get_settings  # noqa: E402


def load_labelled(path: Path) -> list[tuple[str | None, str | None, str]]:
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    ia, ib, il = header.index("addressa"), header.index("addressb"), header.index("truelabel")
    out = []
    for r in rows:
        if all(v is None for v in r):
            continue
        label = str(r[il]).strip().upper() if r[il] is not None else "UNKNOWN"
        out.append((r[ia], r[ib], label))
    wb.close()
    return out


def evaluate(rows, normalizer: CachingNormalizer, cfg: ScoringConfig) -> dict[str, float]:
    m = AddressMatcher(normalizer, cfg)
    tally = {k: 0 for k in ("tp", "fp", "fn_review", "fnm", "tn", "review_diff", "n_same", "n_diff", "review")}
    for a, b, label in rows:
        if label not in ("SAME", "DIFFERENT"):
            continue
        d = m.match(a, b).decision
        if label == "SAME":
            tally["n_same"] += 1
            tally["tp"] += d == "MATCH"
            tally["fnm"] += d == "NON-MATCH"
            tally["fn_review"] += d == "REVIEW"
        else:
            tally["n_diff"] += 1
            tally["fp"] += d == "MATCH"
            tally["tn"] += d == "NON-MATCH"
            tally["review_diff"] += d == "REVIEW"
        tally["review"] += d == "REVIEW"
    labelled = tally["n_same"] + tally["n_diff"]
    div = lambda x, y: x / y if y else float("nan")  # noqa: E731
    return {
        "labelled_rows": labelled,
        "precision": div(tally["tp"], tally["tp"] + tally["fp"]),
        "recall": div(tally["tp"], tally["n_same"]),
        "false_match_rate": div(tally["fp"], tally["n_diff"]),
        "false_non_match_rate": div(tally["fnm"], tally["n_same"]),
        "review_rate": div(tally["review"], labelled),
    }


def fmt(metrics: dict[str, float]) -> str:
    return "  ".join(f"{k}={v:.2f}" if isinstance(v, float) else f"{k}={v}" for k, v in metrics.items())


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("labelled_xlsx", type=Path)
    ap.add_argument("--sweep", action="store_true", help="try a grid of match/non-match thresholds")
    args = ap.parse_args()
    settings = get_settings()
    cfg = load_scoring_config(settings.scoring_config_path)
    normalizer = CachingNormalizer(get_libpostal_backend(settings.libpostal_languages),
                                   max_chars=settings.max_address_chars, max_expansions=cfg.max_expansions)
    rows = load_labelled(args.labelled_xlsx)
    print("current config:", fmt(evaluate(rows, normalizer, cfg)))
    if args.sweep:
        print("\nmatch  non_match  metrics")
        for mt in range(70, 100, 5):
            for nt in range(30, min(mt, 70) + 1, 10):
                c = dataclasses.replace(cfg, match_threshold=mt, non_match_threshold=nt)
                print(f"{mt:>5}  {nt:>9}  {fmt(evaluate(rows, normalizer, c))}")
        print("\nPick thresholds on a validation split, confirm on a held-out test split.")


if __name__ == "__main__":
    main()
