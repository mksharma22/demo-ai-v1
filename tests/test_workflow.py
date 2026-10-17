"""End-to-end orchestration through the LangGraph workflow with a fake backend."""

from __future__ import annotations

import pytest
from conftest import FakeBackend, write_xlsx
from openpyxl import load_workbook

from addrmatch.excel_io import OUTPUT_COLUMNS
from addrmatch.jobs import JobManager

ROWS = [
    ["12 Main St, Springfield, IL 62701", "12 main street, springfield, il 62701"],
    ["12 Main St, Springfield, IL 62701", "21 Main St, Springfield, IL 62701"],
    [None, "12 Main St, Springfield, IL 62701"],
    [None, None],
    ["BOOM", "12 Main St, Springfield, IL 62701"],
    ["12 Main St Apt 4b, Springfield, IL 62701", "12 Main St Apt 5c, Springfield, IL 62701"],
    ["12 Main St, Springfield, IL 62701", "12 main street, springfield, il 62701"],  # duplicate
    ["=1+1", "x"],
]


def submit(manager, tmp_path, rows=ROWS, header=("AddressA", "AddressB", "Extra")):
    p = write_xlsx(tmp_path / "in.xlsx", list(header), [r + ["ignored"] for r in rows])
    job_id, reused = manager.submit("in.xlsx", p.read_bytes())
    return job_id, reused, manager.wait(job_id, timeout=60)


@pytest.fixture
def manager(settings, cfg):
    m = JobManager(settings, backend=FakeBackend(fail_on={"BOOM"}), scoring=cfg)
    yield m
    m.shutdown()


def test_end_to_end_output_schema_order_and_counts(manager, tmp_path):
    job_id, reused, meta = submit(manager, tmp_path)
    assert not reused
    assert meta.status == "completed", meta.error_message
    store = manager.store(job_id)
    ws = load_workbook(store.output_path).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == OUTPUT_COLUMNS
    body = rows[1:]
    assert len(body) == len(ROWS)  # one output row per input row, none dropped
    assert [r[0] for r in body] == [r[0] for r in ROWS]  # order + original text preserved
    assert [r[1] for r in body] == [r[1] for r in ROWS]
    results = [r[4] for r in body]
    assert results[0].startswith("MATCH | Score: 100/100")
    assert results[1].startswith("NON-MATCH") and "12 vs 21" in results[1]
    assert results[2] == "REVIEW | Score: N/A | AddressA is empty; unable to compare."
    assert results[3] == "REVIEW | Score: N/A | AddressA and AddressB are empty; unable to compare."
    assert results[4].startswith("REVIEW | Score: N/A | Normalization failed for AddressA")
    assert results[5].startswith("NON-MATCH") and "Unit differs" in results[5]
    assert results[6] == results[0]
    assert ws.cell(row=9, column=1).data_type == "s"  # "=1+1" stays text
    assert meta.metrics["duplicate_rows"] == 1
    assert meta.metrics["normalization_failures"] == 1
    assert sum(meta.metrics["decisions"].values()) == len(ROWS)
    assert not store.input_path.exists()  # upload removed after completion


def test_duplicate_submission_reuses_job(manager, tmp_path):
    content = write_xlsx(tmp_path / "in.xlsx", ["AddressA", "AddressB"], ROWS).read_bytes()
    first, reused_first = manager.submit("in.xlsx", content)
    second, reused_second = manager.submit("copy.xlsx", content)  # same bytes, even while running
    manager.wait(first, timeout=60)
    assert not reused_first and reused_second and first == second
    assert len(manager.list_jobs()) == 1


def test_missing_columns_fail_cleanly(manager, tmp_path):
    _, _, meta = submit(manager, tmp_path, header=("Foo", "Bar", "Baz"))
    assert meta.status == "failed"
    assert "AddressA and AddressB were not found" in meta.error_message
    assert "Traceback" not in meta.error_message and "/" not in meta.error_message


def test_malformed_file_fails_cleanly(manager):
    job_id, _ = manager.submit("bad.xlsx", b"not an excel file")
    meta = manager.wait(job_id, timeout=30)
    assert meta.status == "failed" and "not a valid .xlsx" in meta.error_message


def test_header_only_file_produces_empty_output(manager, tmp_path):
    job_id, _, meta = submit(manager, tmp_path, rows=[])
    assert meta.status == "completed"
    rows = list(load_workbook(manager.store(job_id).output_path).active.iter_rows(values_only=True))
    assert rows == [OUTPUT_COLUMNS]


def test_crashed_job_fails_safely_and_retry_resumes(manager, tmp_path, monkeypatch):
    from addrmatch import matcher as mm
    from addrmatch.jobs import GENERIC_FAILURE

    calls = {"n": 0}
    original = mm.AddressMatcher.safe_match

    def flaky(self, a, b):
        calls["n"] += 1
        if calls["n"] == 5:  # inside the second batch (batch_size=3)
            raise RuntimeError("simulated worker crash")
        return original(self, a, b)

    monkeypatch.setattr(mm.AddressMatcher, "safe_match", flaky)
    job_id, _, meta = submit(manager, tmp_path)
    assert meta.status == "failed"
    assert meta.error_message == GENERIC_FAILURE  # no stack trace or paths for users
    store = manager.store(job_id)
    assert store.result_count() == 3  # first batch was persisted
    assert not store.output_path.exists()  # no partial output is downloadable

    manager.retry(job_id)
    meta = manager.wait(job_id, timeout=30)
    assert meta.status == "completed"
    assert store.result_count() == len(ROWS)
    rows = list(load_workbook(store.output_path).active.iter_rows(values_only=True))
    assert len(rows) == len(ROWS) + 1


def test_progress_reaches_total(manager, tmp_path):
    _, _, meta = submit(manager, tmp_path)
    assert meta.processed_rows == meta.total_rows == len(ROWS)
    assert meta.progress == 1.0
