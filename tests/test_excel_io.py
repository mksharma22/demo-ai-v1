from __future__ import annotations

import pytest
from conftest import write_xlsx
from openpyxl import load_workbook

from addrmatch.excel_io import (
    OUTPUT_COLUMNS,
    InputValidationError,
    iter_input_rows,
    validate_workbook,
    write_output,
)


def test_missing_required_column(tmp_path):
    p = write_xlsx(tmp_path / "x.xlsx", ["AddressA", "Other"], [["a", "b"]])
    with pytest.raises(InputValidationError, match="AddressA and AddressB were not found"):
        validate_workbook(p)


def test_headers_are_case_and_space_insensitive_and_extras_ignored(tmp_path):
    p = write_xlsx(tmp_path / "x.xlsx", ["ID", " addressa ", "ADDRESSB", "Normalized Address"], [[1, "a", "b", "x"]])
    info = validate_workbook(p)
    assert (info.col_a, info.col_b) == (1, 2)
    assert "Normalized Address" in info.extra_columns


def test_duplicate_required_header(tmp_path):
    p = write_xlsx(tmp_path / "x.xlsx", ["AddressA", "AddressA", "AddressB"], [])
    with pytest.raises(InputValidationError, match="duplicate"):
        validate_workbook(p)


def test_wrong_extension(tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("AddressA,AddressB\n")
    with pytest.raises(InputValidationError, match=".xlsx"):
        validate_workbook(p)


def test_not_a_zip(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"this is not a workbook")
    with pytest.raises(InputValidationError, match="not a valid .xlsx"):
        validate_workbook(p)


def test_empty_file(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"")
    with pytest.raises(InputValidationError, match="empty"):
        validate_workbook(p)


def test_zip_without_workbook_part(tmp_path):
    import zipfile

    p = tmp_path / "x.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(InputValidationError, match="missing workbook part"):
        validate_workbook(p)


def test_size_limit(tmp_path):
    p = write_xlsx(tmp_path / "x.xlsx", ["AddressA", "AddressB"], [["a" * 1000, "b"]] * 50)
    with pytest.raises(InputValidationError, match="larger than"):
        validate_workbook(p, max_upload_mb=0)


def test_rows_preserved_blank_interior_kept_trailing_dropped(tmp_path):
    rows = [["12 Main St", "x"], [None, None], ["   ", "13 Main St"], [12, 3.5], [None, None], [None, None]]
    p = write_xlsx(tmp_path / "x.xlsx", ["AddressA", "AddressB"], rows)
    info = validate_workbook(p)
    out = list(iter_input_rows(p, info, max_rows=100))
    assert [r.address_a for r in out] == ["12 Main St", None, "   ", "12"]
    assert [r.address_b for r in out] == ["x", None, "13 Main St", "3.5"]
    assert [r.excel_row for r in out] == [2, 3, 4, 5]
    assert [r.row_index for r in out] == [0, 1, 2, 3]


def test_formula_cells_read_as_literal_text(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["AddressA", "AddressB"])
    ws.append(["=1+1", "12 Main St"])  # a real formula cell
    p = tmp_path / "f.xlsx"
    wb.save(p)
    info = validate_workbook(p)
    (row,) = list(iter_input_rows(p, info, max_rows=10))
    assert row.address_a == "=1+1"


def test_max_rows(tmp_path):
    p = write_xlsx(tmp_path / "x.xlsx", ["AddressA", "AddressB"], [["a", "b"]] * 5)
    info = validate_workbook(p)
    with pytest.raises(InputValidationError, match="more than 3"):
        list(iter_input_rows(p, info, max_rows=3))


def test_output_schema_and_formula_safety(tmp_path):
    p = tmp_path / "out.xlsx"
    n = write_output([("=HYPERLINK(\"x\")", "b", "na", "nb", "MATCH | Score: 100/100 | ok.")], p)
    assert n == 1
    ws = load_workbook(p).active
    assert tuple(c.value for c in ws[1]) == OUTPUT_COLUMNS
    cell = ws.cell(row=2, column=1)
    assert cell.value == '=HYPERLINK("x")' and cell.data_type == "s"


def test_output_rejects_wrong_width(tmp_path):
    with pytest.raises(ValueError):
        write_output([("a", "b")], tmp_path / "o.xlsx")
