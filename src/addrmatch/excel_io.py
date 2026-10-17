"""Safe XLSX reading (streaming, no formula/macro execution) and writing."""

from __future__ import annotations

import contextlib
import datetime as dt
import io
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import IO, Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

REQUIRED_COLUMNS = ("AddressA", "AddressB")
OUTPUT_COLUMNS = ("AddressA", "AddressB", "normalized_AddressA", "normalized_AddressB", "Result")
EXCEL_MAX_CELL_CHARS = 32_767


class InputValidationError(ValueError):
    """User-facing problem with the uploaded workbook. The message is safe to display."""


@dataclass(frozen=True)
class WorkbookInfo:
    sheet_name: str
    col_a: int  # 0-based column index
    col_b: int
    extra_columns: tuple[str, ...]


@dataclass(frozen=True)
class InputRow:
    row_index: int     # 0-based position among data rows (output order)
    excel_row: int     # 1-based row number in the sheet, for user messages
    address_a: str | None
    address_b: str | None


def check_file(path: Path, *, max_upload_mb: int, max_uncompressed_mb: int) -> None:
    """Cheap structural checks before openpyxl touches the file."""
    if path.suffix.lower() != ".xlsx":
        raise InputValidationError("Only .xlsx files are accepted (not .xls, .xlsm or .csv).")
    size = path.stat().st_size
    if size == 0:
        raise InputValidationError("The uploaded file is empty.")
    if size > max_upload_mb * 1024 * 1024:
        raise InputValidationError(f"File is larger than the {max_upload_mb} MB limit.")
    if not zipfile.is_zipfile(path):
        raise InputValidationError("The file is not a valid .xlsx workbook (it is not a zip package).")
    try:
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            total = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise InputValidationError("The workbook appears to be corrupted.") from exc
    if "xl/workbook.xml" not in names:
        raise InputValidationError("The file is not a valid .xlsx workbook (missing workbook part).")
    if total > max_uncompressed_mb * 1024 * 1024:
        raise InputValidationError("The workbook expands to an unexpectedly large size and was rejected.")


def cell_to_text(value: Any) -> str | None:
    """Represent a cell value as text without altering string content."""
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _open(path: Path):
    # read_only streams rows; data_only=False returns formula text instead of
    # cached results, so a formula is treated as literal text. openpyxl never
    # evaluates formulas or runs VBA (keep_vba defaults to False).
    try:
        return load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise InputValidationError("The workbook could not be opened; re-save it as .xlsx and retry.") from exc


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, v in enumerate(row):
        if v is None:
            continue
        key = str(v).strip().lower()
        if key:
            out.setdefault(key, i)
    return out


def validate_workbook(path: Path, *, max_upload_mb: int = 20, max_uncompressed_mb: int = 200) -> WorkbookInfo:
    """Find the first sheet whose header row has AddressA and AddressB (case-insensitive)."""
    check_file(path, max_upload_mb=max_upload_mb, max_uncompressed_mb=max_uncompressed_mb)
    wb = _open(path)
    try:
        seen: list[str] = []
        for ws in wb.worksheets:
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            hmap = _header_map(header)
            seen.append(ws.title)
            names = [str(h).strip().lower() for h in header if h is not None and str(h).strip()]
            dupes = {c for c in ("addressa", "addressb") if names.count(c) > 1}
            if dupes:
                raise InputValidationError(f"Sheet '{ws.title}' has duplicate column(s): {', '.join(sorted(dupes))}.")
            if "addressa" in hmap and "addressb" in hmap:
                extras = tuple(
                    str(h).strip() for h in header
                    if h is not None and str(h).strip().lower() not in ("addressa", "addressb")
                )
                return WorkbookInfo(ws.title, hmap["addressa"], hmap["addressb"], extras)
        raise InputValidationError(
            "Required columns AddressA and AddressB were not found in the first row of any sheet "
            f"(sheets checked: {', '.join(seen) or 'none'})."
        )
    finally:
        wb.close()


def iter_input_rows(path: Path, info: WorkbookInfo, *, max_rows: int) -> Iterator[InputRow]:
    """Stream data rows in sheet order. Blank rows inside the data are kept;
    fully blank rows after the last non-blank row (formatting residue) are not emitted."""
    wb = _open(path)
    try:
        ws = wb[info.sheet_name]
        pending_blank: list[int] = []
        index = 0
        for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            a = cell_to_text(row[info.col_a]) if info.col_a < len(row) else None
            b = cell_to_text(row[info.col_b]) if info.col_b < len(row) else None
            if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row):
                pending_blank.append(excel_row)
                continue
            for blank_row in pending_blank:  # interior blank rows are real input rows
                yield _guard(InputRow(index, blank_row, None, None), max_rows)
                index += 1
            pending_blank.clear()
            yield _guard(InputRow(index, excel_row, a, b), max_rows)
            index += 1
    finally:
        wb.close()


def _guard(row: InputRow, max_rows: int) -> InputRow:
    if row.row_index >= max_rows:
        raise InputValidationError(f"The sheet has more than {max_rows:,} data rows (POC limit).")
    return row


def _text_cell(ws, value: str | None, bold: bool = False) -> WriteOnlyCell:
    text = "" if value is None else ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_MAX_CELL_CHARS]
    cell = WriteOnlyCell(ws, value=text)
    cell.data_type = "s"  # force literal text: "=..." is never written as a formula
    if bold:
        cell.font = Font(bold=True)
    else:
        cell.alignment = Alignment(wrap_text=False, vertical="top")
    return cell


def write_output(rows: Iterable[tuple[str | None, ...]], target: Path | IO[bytes]) -> int:
    """Write the five-column output workbook. Returns the number of data rows."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Results")
    for letter, width in zip("ABCDE", (45, 45, 50, 50, 90)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.append([_text_cell(ws, c, bold=True) for c in OUTPUT_COLUMNS])
    n = 0
    try:
        for row in rows:
            if len(row) != len(OUTPUT_COLUMNS):
                raise ValueError("Output rows must have exactly five values")
            ws.append([_text_cell(ws, v) for v in row])
            n += 1
    except BaseException:
        with contextlib.suppress(Exception):
            ws.close()  # release openpyxl's streaming writer before propagating
        raise
    wb.save(target)
    return n


def output_bytes(rows: Iterable[tuple[str | None, ...]]) -> tuple[bytes, int]:
    buf = io.BytesIO()
    n = write_output(rows, buf)
    return buf.getvalue(), n
